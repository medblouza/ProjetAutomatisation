import math
import streamlit as st
import requests
from io import BytesIO
from docx import Document


def generate_docx(content: str) -> bytes:
    doc = Document()
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("**"))
            run.bold = True
        else:
            doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ─── Config ─────────────────────────────────────────
st.set_page_config(page_title="DP Explorer", page_icon="🔍")

API_BASE = "http://127.0.0.1:8000"

# ─── Session ────────────────────────────────────────
if "auth" not in st.session_state:
    st.session_state.auth = None
if "result" not in st.session_state:
    st.session_state.result = None
if "cdc" not in st.session_state:
    st.session_state.cdc = None
if "cleaned" not in st.session_state:
    st.session_state.cleaned = None
if "cleaned_data" not in st.session_state:       # ← pour le wireframe
    st.session_state.cleaned_data = None

# Dropbox
if "dbx_access_token" not in st.session_state:
    st.session_state.dbx_access_token = ""
if "dbx_files" not in st.session_state:
    st.session_state.dbx_files = []
if "dbx_current_path" not in st.session_state:
    st.session_state.dbx_current_path = ""
if "dbx_path_history" not in st.session_state:
    st.session_state.dbx_path_history = []
if "dbx_auth_url" not in st.session_state:
    st.session_state.dbx_auth_url = ""
if "dbx_auth_state" not in st.session_state:
    st.session_state.dbx_auth_state = ""
if "dbx_check_report" not in st.session_state:
    st.session_state.dbx_check_report = None

# ─── Sidebar ────────────────────────────────────────
st.sidebar.title("🔧 Configuration")

username = st.sidebar.text_input("Username")
password = st.sidebar.text_input("Password", type="password")

if st.sidebar.button("Connexion"):
    try:
        res = requests.post(
            f"{API_BASE}/login",
            json={"username": username, "password": password},
            timeout=10,
        )
        if res.status_code == 200:
            st.session_state.auth = {"username": username, "password": password}
            st.sidebar.success("Connecté")
        else:
            st.sidebar.error("Erreur login")
    except Exception as e:
        st.sidebar.error(f"Erreur API : {e}")

if st.session_state.auth:
    if st.sidebar.button("Déconnexion"):
        st.session_state.auth = None
        st.session_state.result = None
        st.session_state.cdc = None
        st.session_state.cleaned = None
        st.session_state.cleaned_data = None      # ← reset wireframe aussi
        st.rerun()

st.sidebar.divider()
if st.session_state.dbx_access_token:
    st.sidebar.success("📦 Dropbox connecté")
    if st.sidebar.button("🔓 Déconnecter Dropbox"):
        st.session_state.dbx_access_token = ""
        st.session_state.dbx_files = []
        st.session_state.dbx_current_path = ""
        st.session_state.dbx_path_history = []
        st.session_state.dbx_auth_url = ""
        st.session_state.dbx_auth_state = ""
        st.session_state.dbx_check_report = None
        st.rerun()
else:
    st.sidebar.info("📦 Dropbox non connecté")

# Indicateur sidebar si données nettoyées disponibles pour le wireframe
if st.session_state.cleaned_data:
    st.sidebar.success("🎨 Données prêtes pour le wireframe")

# ─── TABS ────────────────────────────────────────────
tab_dp, tab_dbx, tab_drive, tab_wireframe, tab_style, tab_sitegen = st.tabs(
    ["📊 DP Explorer", "📦 Dropbox", "📁 Dossier partagé", "🎨 Wireframe", "🖌️ Style Extractor", "🏗️ Site Generator"]
)


# ════════════════════════════════════════════════════
# TAB 5 — Style Extractor
# ════════════════════════════════════════════════════
with tab_style:
    try:
        from streamlit_style_tab import render_style_tab
        render_style_tab()
    except Exception as e:
        st.error(f"Erreur chargement Style Extractor : {e}")

# ════════════════════════════════════════════════════
# TAB 6 — Site Gen
# ════════════════════════════════════════════════════
with tab_sitegen:
    try:
        from streamlit_sitegen_tab import render_sitegen_tab
        render_sitegen_tab()
    except Exception as e:
        st.error(f"Erreur chargement Site Gen : {e}")

# ════════════════════════════════════════════════════
# TAB 1 — DP Explorer
# ════════════════════════════════════════════════════
with tab_dp:

    st.title("📊 DP Data Explorer")
    code = st.text_input("Code client (Code Sage)")
    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("🔍 Récupérer données"):
            if not st.session_state.auth:
                st.error("Connecte-toi d'abord")
            elif not code:
                st.warning("Entre un code")
            else:
                with st.spinner("Chargement..."):
                    try:
                        res = requests.post(
                            f"{API_BASE}/info",
                            json={
                                "username": st.session_state.auth["username"],
                                "password": st.session_state.auth["password"],
                                "code": code,
                            },
                            timeout=60,
                        )
                        if res.status_code == 200:
                            st.session_state.result = res.json()
                        else:
                            st.error(res.json().get("detail", "Erreur API"))
                    except Exception as e:
                        st.error(f"Erreur connexion API : {e}")

    with col2:
        if st.button("🧾 Générer Cahier de charge"):
            if not st.session_state.auth:
                st.error("Connecte-toi d'abord")
            elif not code:
                st.warning("Entre un code")
            else:
                with st.spinner("Génération du cahier de charge..."):
                    try:
                        res = requests.post(
                            f"{API_BASE}/generate-cdc",
                            json={
                                "username": st.session_state.auth["username"],
                                "password": st.session_state.auth["password"],
                                "code": code,
                            },
                            timeout=600,
                        )
                        if res.status_code == 200:
                            data = res.json()
                            st.session_state.cdc = data.get("content")
                            st.session_state.qa  = {"score": data.get("score")}
                        else:
                            st.error(res.json().get("detail", "Erreur API"))
                    except Exception as e:
                        st.error(f"Erreur connexion API : {e}")

    with col3:
        if st.button("🧹 Nettoyer données"):
            if not st.session_state.auth:
                st.error("Connecte-toi d'abord")
            elif not code:
                st.warning("Entre un code")
            else:
                with st.spinner("Nettoyage..."):
                    try:
                        res = requests.post(
                            f"{API_BASE}/clean",
                            json={
                                "username": st.session_state.auth["username"],
                                "password": st.session_state.auth["password"],
                                "code": code,
                            },
                            timeout=60,
                        )
                        if res.status_code == 200:
                            cleaned_result = res.json().get("cleaned")
                            st.session_state.cleaned = cleaned_result       # affichage JSON
                            st.session_state.cleaned_data = cleaned_result  # ← wireframe
                            st.success("✅ Données nettoyées — tu peux générer le wireframe dans l'onglet 🎨")
                        else:
                            st.error(res.json().get("detail", "Erreur API"))
                    except Exception as e:
                        st.error(f"Erreur : {e}")

    if st.session_state.result:
        data = st.session_state.result
        st.subheader("📌 Résumé")
        c1, c2, c3 = st.columns(3)
        with c1:
            st.metric("Enseigne", data.get("company", {}).get("Name", "—"))
        with c2:
            st.metric("Secteur", data.get("company", {}).get("Industry", "—"))
        with c3:
            st.metric("Ville", data.get("company", {}).get("BillingCity", "—"))
        st.divider()
        st.subheader("📦 Données complètes")
        with st.expander("Partner"):
            st.json(data.get("partner", {}))
        with st.expander("Details"):
            st.json(data.get("details", {}))
        with st.expander("Company"):
            st.json(data.get("company", {}))

    if st.session_state.cleaned:
        st.divider()
        st.subheader("🧹 Données nettoyées")
        st.json(st.session_state.cleaned)

    if st.session_state.cdc:
        st.divider()
        st.subheader("📄 Cahier de charge généré")
        if "qa" in st.session_state:
            st.info(f"Score qualité : {st.session_state.qa.get('score', 'N/A')}/100")
        st.text_area("CDC", st.session_state.cdc, height=400)
        st.download_button(
            label="📥 Télécharger le CDC",
            data=generate_docx(st.session_state.cdc),
            file_name="cahier_de_charge.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


# ════════════════════════════════════════════════════
# TAB 2 — Dropbox
# ════════════════════════════════════════════════════
with tab_dbx:

    st.title("📦 Dropbox")

    # ── Helpers ──────────────────────────────────────
    def fmt_size(n: int) -> str:
        if not n:
            return "—"
        units = ["o", "Ko", "Mo", "Go"]
        i = min(int(math.log(n, 1024)), len(units) - 1)
        return f"{n / (1024 ** i):.1f} {units[i]}"

    def file_icon(f: dict) -> str:
        if f.get("is_folder"):
            return "📁"
        name = f.get("name", "").lower()
        if name.endswith(".pdf"):                                      return "📄"
        if name.endswith((".jpg", ".jpeg", ".png", ".gif", ".webp")): return "🖼️"
        if name.endswith((".mp4", ".mov")):                            return "🎬"
        if name.endswith((".mp3", ".wav")):                            return "🎵"
        if name.endswith((".zip", ".tar", ".gz")):                     return "🗜️"
        if name.endswith((".xlsx", ".xls", ".csv")):                   return "📊"
        if name.endswith((".docx", ".doc")):                           return "📝"
        if name.endswith((".py", ".js", ".json")):                     return "💻"
        return "📎"

    def load_files(path: str):
        with st.spinner("Chargement..."):
            try:
                res = requests.post(
                    f"{API_BASE}/dropbox/list",
                    json={
                        "access_token": st.session_state.dbx_access_token,
                        "path": path,
                        "recursive": False,
                    },
                    timeout=30,
                )
                if res.status_code == 200:
                    st.session_state.dbx_files        = res.json().get("files", [])
                    st.session_state.dbx_current_path = path
                else:
                    st.error(res.json().get("detail", "Erreur API"))
            except Exception as e:
                st.error(f"Erreur : {e}")

    def search_files(query: str):
        with st.spinner("Recherche..."):
            try:
                res = requests.post(
                    f"{API_BASE}/dropbox/search",
                    json={
                        "access_token": st.session_state.dbx_access_token,
                        "query": query,
                        "path": "",
                    },
                    timeout=30,
                )
                if res.status_code == 200:
                    data = res.json()
                    st.session_state.dbx_files = data.get("files", [])
                    st.info(f"{data.get('total', 0)} résultat(s) pour « {query} »")
                else:
                    st.error(res.json().get("detail", "Erreur API"))
            except Exception as e:
                st.error(f"Erreur : {e}")

    # ── Authentification ─────────────────────────────
    if not st.session_state.dbx_access_token:
        st.info("Connecte-toi à Dropbox pour accéder à tes fichiers.")
        col_a, col_b = st.columns([2, 1])

        with col_a:
            if st.button("🔗 Se connecter à Dropbox", use_container_width=True):
                try:
                    res = requests.get(f"{API_BASE}/dropbox/auth-url", timeout=10)
                    if res.status_code == 200:
                        data = res.json()
                        st.session_state.dbx_auth_url   = data["auth_url"]
                        st.session_state.dbx_auth_state = data["state"]
                    else:
                        st.error("Impossible d'obtenir l'URL d'auth.")
                except Exception as e:
                    st.error(f"Erreur API : {e}")

            if st.session_state.dbx_auth_url:
                st.markdown(f"**1.** [Clique ici pour autoriser l'accès Dropbox]({st.session_state.dbx_auth_url})")
                st.markdown("**2.** Copie le code affiché après l'autorisation et colle-le ci-dessous :")
                auth_code = st.text_input("Code d'autorisation", placeholder="sl.AbCdEf...")
                if st.button("✅ Valider le code", use_container_width=True):
                    if not auth_code.strip():
                        st.warning("Entre le code d'autorisation.")
                    else:
                        try:
                            res = requests.post(
                                f"{API_BASE}/dropbox/callback",
                                json={"code": auth_code.strip(), "state": st.session_state.dbx_auth_state},
                                timeout=15,
                            )
                            if res.status_code == 200:
                                st.session_state.dbx_access_token = res.json().get("access_token", "")
                                st.success("✅ Connecté à Dropbox !")
                                st.rerun()
                            else:
                                st.error(res.json().get("detail", "Erreur d'authentification."))
                        except Exception as e:
                            st.error(f"Erreur : {e}")

        with col_b:
            st.markdown("#### Comment ça marche ?")
            st.markdown("1. Clique **Se connecter**\n2. Autorise l'app\n3. Copie le **code**\n4. Colle-le et valide")

    # ── Interface fichiers ────────────────────────────
    else:
        st.success("✅ Connecté à Dropbox")
        st.divider()

        nav_col, search_col = st.columns([3, 2])

        with nav_col:
            breadcrumb = " › ".join(["🏠 Racine"] + st.session_state.dbx_path_history)
            st.caption(f"📂 {breadcrumb}")
            path_input = st.text_input(
                "Chemin (vide = racine)",
                value=st.session_state.dbx_current_path,
                placeholder="/Mon Dossier",
            )

            c1, c2, c3 = st.columns(3)
            with c1:
                if st.button("📂 Lister", use_container_width=True):
                    load_files(path_input.strip())
                    st.session_state.dbx_check_report = None
            with c2:
                if st.button("🔍 Auto-check assets", use_container_width=True):
                    with st.spinner("Analyse en cours... (quelques secondes par fichier)"):
                        try:
                            res = requests.post(
                                f"{API_BASE}/dropbox/check-assets",
                                json={
                                    "access_token": st.session_state.dbx_access_token,
                                    "path": path_input.strip(),
                                },
                                timeout=120,
                            )
                            if res.status_code == 200:
                                st.session_state.dbx_check_report = res.json()
                                st.success("Analyse terminée !")
                            else:
                                st.error(res.json().get("detail", "Erreur API"))
                        except Exception as e:
                            st.error(f"Erreur : {e}")
            with c3:
                if st.session_state.dbx_path_history:
                    if st.button("⬅️ Retour", use_container_width=True):
                        st.session_state.dbx_path_history.pop()
                        parent = "/".join(
                            st.session_state.dbx_current_path.rsplit("/", 1)[:-1]
                        ) or ""
                        load_files(parent)

        with search_col:
            search_query = st.text_input("🔍 Recherche", placeholder="rapport 2024...")
            if st.button("Rechercher", use_container_width=True):
                if search_query.strip():
                    search_files(search_query.strip())
                else:
                    st.warning("Entre un terme de recherche.")

        st.divider()

        # ── Rapport Auto-check ────────────────────────
        if st.session_state.dbx_check_report:
            report = st.session_state.dbx_check_report
            st.subheader("🔍 Rapport Auto-check Assets")

            m1, m2, m3, m4, m5 = st.columns(5)
            m1.metric("Total",         report.get("total", 0))
            m2.metric("✅ Validés",    report.get("ok", 0))
            m3.metric("⚠️ Attention", report.get("warnings", 0))
            m4.metric("⛔ Bloqués",   report.get("errors", 0))
            m5.metric("❓ À classer", report.get("needs_manual", 0))

            st.divider()

            rf1, rf2 = st.columns(2)
            with rf1:
                filtre_statut = st.selectbox(
                    "Statut",
                    ["Tout", "✅ Validés", "⚠️ Attention", "⛔ Bloqués", "❓ À classer"],
                    key="filter_statut",
                )
            with rf2:
                filtre_cat = st.selectbox(
                    "Catégorie",
                    ["Tout", "logo", "icon", "image", "banner", "video", "svg", "doc", "unknown"],
                    key="filter_cat",
                )

            assets = report.get("assets", [])
            if filtre_statut == "✅ Validés":
                assets = [a for a in assets if a["status"] == "ok"]
            elif filtre_statut == "⚠️ Attention":
                assets = [a for a in assets if a["status"] == "warning"]
            elif filtre_statut == "⛔ Bloqués":
                assets = [a for a in assets if a["status"] == "error"]
            elif filtre_statut == "❓ À classer":
                assets = [a for a in assets if a.get("needs_manual")]
            if filtre_cat != "Tout":
                assets = [a for a in assets if a.get("category") == filtre_cat]

            cat_icons = {
                "logo": "🔤", "icon": "🔷", "image": "🖼️", "banner": "🖼️",
                "video": "🎬", "svg": "✏️", "doc": "📄", "unknown": "❓",
            }

            if not assets:
                st.info("Aucun asset ne correspond aux filtres.")
            else:
                for asset in assets:
                    badge    = asset.get("badge", "")
                    name     = asset.get("name", "")
                    ext      = asset.get("extension", "—")
                    size     = asset.get("size", 0)
                    category = asset.get("category", "unknown")
                    label    = asset.get("label", "")
                    dims     = asset.get("dimensions")
                    issues   = asset.get("issues", [])
                    sugs     = asset.get("suggestions", [])
                    manual   = asset.get("needs_manual", False)
                    auto     = asset.get("category_auto", False)

                    size_str = (
                        f"{size / (1024*1024):.1f} Mo" if size > 1024*1024
                        else f"{size / 1024:.1f} Ko" if size > 1024
                        else f"{size} o"
                    )
                    dims_str  = f"{dims[0]}x{dims[1]}px" if dims else "—"
                    cat_icon  = cat_icons.get(category, "❓")
                    cat_label = f"{cat_icon} {category}" + (" ❓" if not auto else "")

                    col1, col2, col3 = st.columns([0.5, 4, 3])
                    with col1:
                        st.markdown(f"## {badge}")
                    with col2:
                        st.markdown(f"**{name}**")
                        st.caption(f"{cat_label} — {ext} — {size_str} — {dims_str}")
                        if manual:
                            new_cat = st.selectbox(
                                "Préciser la catégorie :",
                                ["logo", "icon", "image", "banner", "video", "doc"],
                                key=f"cat_{asset['path']}",
                            )
                            if st.button("✅ Confirmer", key=f"confirm_{asset['path']}"):
                                st.success(f"Catégorie : {new_cat}")
                    with col3:
                        if issues:
                            for issue in issues:
                                st.warning(issue)
                            st.divider()
                            for sug in sugs:
                                st.caption(f"💡 {sug}")
                        else:
                            st.success(f"✅ {label} — Conforme")
                    st.divider()

            # ═══════════════════════════════════════════════
                # UI CSS
            # ═══════════════════════════════════════════════

    st.markdown("""
    <style>

    .file-card {
        background-color: #111827;
        border: 1px solid #2d3748;
        border-radius: 18px;
        padding: 12px;
        transition: 0.2s ease-in-out;
        min-height: 340px;
        margin-bottom: 18px;
    }

    .file-card:hover {
        border-color: #6366f1;
        transform: translateY(-2px);
    }

    .file-title {
        font-size: 14px;
        font-weight: 600;
        margin-top: 10px;
        word-break: break-word;
    }

    .file-meta {
        color: #9ca3af;
        font-size: 12px;
    }

    .folder-box {
        display:flex;
        align-items:center;
        justify-content:center;
        height:180px;
        font-size:70px;
    }

    .file-box {
        display:flex;
        align-items:center;
        justify-content:center;
        height:180px;
        font-size:55px;
    }

    div[data-testid="stImage"] img {
        border-radius: 14px;
    }

    button[kind="secondary"] {
        border-radius: 10px;
    }

    </style>
    """, unsafe_allow_html=True)

    # ═══════════════════════════════════════════════
    # LISTE DES FICHIERS
    # ═══════════════════════════════════════════════

    files = st.session_state.dbx_files

    if files:

        st.subheader(f"📂 {len(files)} élément(s)")

        # ─────────────────────────────────────
        # FILTRES
        # ─────────────────────────────────────

        f1, f2 = st.columns([1, 2])

        with f1:
            filter_type = st.selectbox(
                "Type",
                ["Tout", "Dossiers", "Fichiers"]
            )

        with f2:
            filter_name = st.text_input(
                "Recherche",
                placeholder="logo, banner..."
            )

        filtered = files

        if filter_type == "Dossiers":
            filtered = [
                f for f in filtered
                if f.get("is_folder")
            ]

        elif filter_type == "Fichiers":
            filtered = [
                f for f in filtered
                if not f.get("is_folder")
            ]

        if filter_name:
            filtered = [
                f for f in filtered
                if filter_name.lower() in f.get("name", "").lower()
            ]

        if not filtered:

            st.warning("Aucun élément trouvé.")

        else:

            cols = st.columns(4)

            image_exts = (
                ".png",
                ".jpg",
                ".jpeg",
                ".gif",
                ".webp"
            )

            for idx, f in enumerate(filtered):

                col = cols[idx % 4]

                with col:

                    name = f.get("name", "")
                    path = f.get("path", "")
                    size = fmt_size(f.get("size", 0))
                    date = (f.get("modified", "") or "")[:10]
                    is_dir = f.get("is_folder", False)

                    with st.container(border=True):

                        # ═══════════════════════════════
                        # DOSSIER
                        # ═══════════════════════════════

                        if is_dir:

                            st.markdown(
                                """
                                <div class="folder-box">
                                    📁
                                </div>
                                """,
                                unsafe_allow_html=True
                            )

                            st.markdown(f"**{name}**")
                            st.caption("Dossier Dropbox")

                            if st.button(
                                "📂 Ouvrir",
                                key=f"open_{path}",
                                use_container_width=True
                            ):

                                st.session_state.dbx_path_history.append(name)
                                load_files(path)

                        # ═══════════════════════════════
                        # IMAGE
                        # ═══════════════════════════════

                        elif name.lower().endswith(image_exts):

                            preview_data = None

                            try:

                                preview_res = requests.post(
                                    f"{API_BASE}/dropbox/download",
                                    json={
                                        "access_token": st.session_state.dbx_access_token,
                                        "path": path,
                                    },
                                    timeout=30,
                                )

                                if preview_res.status_code == 200:
                                    preview_data = preview_res.content

                            except:
                                pass

                            if preview_data:

                                st.image(
                                    preview_data,
                                    use_container_width=True
                                )

                            else:

                                st.markdown(
                                    """
                                    <div class="file-box">
                                        🖼️
                                    </div>
                                    """,
                                    unsafe_allow_html=True
                                )

                            st.markdown(f"**{name}**")
                            st.caption(f"{size} • {date}")

                            if preview_data:

                                st.download_button(
                                    label="⬇️ Télécharger",
                                    data=preview_data,
                                    file_name=name,
                                    use_container_width=True,
                                    key=f"img_dl_{path}"
                                )

                        # ═══════════════════════════════
                        # AUTRES FICHIERS
                        # ═══════════════════════════════

                        else:

                            icon = file_icon(f)

                            st.markdown(
                                f"""
                                <div class="file-box">
                                    {icon}
                                </div>
                                """,
                                unsafe_allow_html=True
                            )

                            st.markdown(f"**{name}**")
                            st.caption(f"{size} • {date}")

                            if st.button(
                                "⬇️ Télécharger",
                                key=f"dl_{path}",
                                use_container_width=True
                            ):

                                with st.spinner("Téléchargement..."):

                                    try:

                                        dl_res = requests.post(
                                            f"{API_BASE}/dropbox/download",
                                            json={
                                                "access_token": st.session_state.dbx_access_token,
                                                "path": path,
                                            },
                                            timeout=60,
                                        )

                                        if dl_res.status_code == 200:

                                            st.download_button(
                                                label="💾 Sauvegarder",
                                                data=dl_res.content,
                                                file_name=name,
                                                mime=dl_res.headers.get(
                                                    "Content-Type",
                                                    "application/octet-stream"
                                                ),
                                                key=f"save_{path}",
                                                use_container_width=True
                                            )

                                        else:

                                            st.error("Erreur téléchargement")

                                    except Exception as e:

                                        st.error(f"Erreur : {e}")


# ════════════════════════════════════════════════════
# TAB 3 — Dossier partagé
# ════════════════════════════════════════════════════
with tab_drive:

    st.title("📁 Dossier partagé")
    st.caption("Accès rapide au dossier partagé client (Google Drive ou autre)")
    st.divider()

    st.subheader("🔗 Lien du dossier")
    col_link, col_btn = st.columns([4, 1])

    with col_link:
        drive_url = st.text_input(
            "URL du dossier partagé",
            placeholder="https://drive.google.com/drive/folders/...",
            label_visibility="collapsed",
        )

    with col_btn:
        if st.button("🌐 Ouvrir", use_container_width=True):
            if drive_url.strip():
                st.markdown(f'<a href="{drive_url}" target="_blank">Cliquez ici pour ouvrir</a>', unsafe_allow_html=True)
            else:
                st.warning("Entre un lien d'abord.")

    if drive_url.strip():
        st.markdown(
            f"""
            <div style="
                background: #f0f2f6;
                border-radius: 8px;
                padding: 16px;
                margin-top: 8px;
                display: flex;
                align-items: center;
                gap: 12px;
            ">
                <span style="font-size: 24px;">📁</span>
                <div>
                    <div style="font-weight: 600;">Dossier partagé client</div>
                    <a href="{drive_url}" target="_blank" style="color: #4a90e2; font-size: 13px;">{drive_url[:60]}{'...' if len(drive_url) > 60 else ''}</a>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        st.markdown("")
        st.link_button("📂 Ouvrir le dossier partagé", drive_url, use_container_width=True)
    else:
        st.info("Aucun lien configuré. Colle le lien du dossier partagé ci-dessus.")

    st.divider()
    st.subheader("📋 À venir")
    st.markdown("""
    - 🔄 Synchronisation automatique avec le dossier client
    - 📊 Liste des fichiers du dossier partagé
    - ✅ Statut des assets depuis le dossier partagé
    - 🔔 Notifications lors d'ajout de nouveaux fichiers
    """)


# ════════════════════════════════════════════════════
# TAB 4 — Wireframe
# ════════════════════════════════════════════════════
with tab_wireframe:
    from streamlit_wireframe_tab import render_wireframe_tab
    render_wireframe_tab()

