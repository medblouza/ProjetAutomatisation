from dotenv import load_dotenv
import os
load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel

from app.social_suggester import suggest_social
from app.dp_service import DpService
from app.cleaner import clean_dp_data
from app.llm.llm_tool import LLMTool
from app.mcp.agents.base_agent import BaseAgent
from app.dropbox_router import dropbox_router
from app.wireframe_router import router as wireframe_router
from app.style_extractor_router import router as style_extractor_router
from app.mcp.routers.design_router import router as design_router
from fastapi.middleware.cors import CORSMiddleware
from app.mcp.routers.web_generator_router import router as web_generator_router
from app.database import init_db, save_cdc
from app.cdc_router import router as cdc_router, design_json_router

app = FastAPI(title="DP Service API")

# Initialisation de la base de données au démarrage
try:
    init_db()
except Exception as _db_exc:
    import logging
    logging.getLogger(__name__).error(f"[DB] Impossible d'initialiser la base : {_db_exc}")

"""app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # DEV ONLY
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)"""

"""app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)"""

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(style_extractor_router, prefix="/api", tags=["Style Extractor"])
app.include_router(dropbox_router, prefix="/dropbox")
app.include_router(wireframe_router, prefix="/api", tags=["Wireframe"])
print(">>> DESIGN ROUTER LOADED <<<")
app.include_router(design_router)
app.include_router(web_generator_router)
app.include_router(cdc_router)
app.include_router(design_json_router)
for route in app.routes:
    print(route.path)
"""app.include_router(design_router)"""
"""app.include_router(site_gen_router, prefix="/api/site-gen", tags=["Site Generator"])"""

# ─────────────────────────────
# Models (inputs API)
# ─────────────────────────────
class AuthRequest(BaseModel):
    username: str
    password: str

class CodeRequest(BaseModel):
    username: str
    password: str
    code: str


# ─────────────────────────────
# Endpoint : test login
# ─────────────────────────────
@app.post("/login")
def login(req: AuthRequest):
    service = DpService(req.username, req.password)

    if service.login():
        return {"status": "success", "message": "Login OK"}
    raise HTTPException(status_code=401, detail="Login failed")


# ─────────────────────────────
# Endpoint : get_info_by_code
# ─────────────────────────────
@app.post("/info")
def get_info(req: CodeRequest):
    service = DpService(req.username, req.password)

    result = service.get_info_by_code(req.code)

    if not result.is_success:
        raise HTTPException(status_code=400, detail=result.message)

    return result.data


# ─────────────────────────────
# Endpoint : process_codesage
# ─────────────────────────────
@app.post("/process")
def process(req: CodeRequest):
    service = DpService(req.username, req.password)

    result = service.process_codesage(req.code)

    if not result.is_success:
        raise HTTPException(status_code=400, detail=result.message)

    return result.data

# ─────────────────────────────
# Endpoint : clean data
# ─────────────────────────────
@app.post("/clean")
def clean(req: CodeRequest):
    service = DpService(req.username, req.password)

    result = service.get_info_by_code(req.code)

    if not result.is_success:
        raise HTTPException(status_code=400, detail=result.message)

    raw_data = result.data

    cleaned = clean_dp_data(raw_data)

    return {
        "cleaned": cleaned
    }

@app.post("/generate-cdc")
def generate_cdc(req: CodeRequest):

    service = DpService(req.username, req.password)
    result  = service.get_info_by_code(req.code)  # ← get_info pas process

    if not result.is_success:
        raise HTTPException(status_code=400, detail=result.message)

    cleaned  = clean_dp_data(result.data)          # ← nettoyage d'abord
    pipeline = BaseAgent(LLMTool())
    output   = pipeline.run(cleaned)               # ← pipeline sur données nettoyées

    content = output.get("content", "")
    score   = output.get("score", 0)

    if isinstance(content, list):
        content = "\n".join([str(x) for x in content])
    else:
        content = str(content)

    try:
        score = int(score)
    except:
        score = 0

    # ── Sauvegarde PostgreSQL (silencieuse — ne bloque jamais la réponse) ──
    save_cdc(company_name=req.code, content=content)

    return {"content": content, "score": score}
@app.post("/suggest-social")
def suggest_social_api(data: dict):

    try:
        result = suggest_social(data)
        return {"suggestions": result}

    except Exception as e:
        return {"error": str(e)}

class DocxRequest(BaseModel):
    content: str
    filename: str = "cahier_de_charge.docx"

@app.post("/api/download-docx")
def download_docx(req: DocxRequest):
    from docx import Document
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    try:
        doc = Document()
        for line in req.content.split("\n"):
            line = line.strip()
            if not line:
                doc.add_paragraph("")
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("**") and line.endswith("**"):
                p = doc.add_paragraph()
                run = p.add_run(line.strip("**"))
                run.bold = True
            else:
                doc.add_paragraph(line)
        
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={req.filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))